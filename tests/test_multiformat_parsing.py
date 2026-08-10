from __future__ import annotations

import importlib.util
import shutil
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas

from app.core.config import get_settings
from app.services.classifier import classify_document
from app.services.extractor import ExtractorService
from app.services.parser import DocumentParser


OCR_FONT = Path("/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf")


def extract_fields(path: Path) -> tuple[object, dict[str, object]]:
    parsed = DocumentParser().parse(path, path.name)
    document_type, _ = classify_document(parsed.text, path.name)
    result = ExtractorService().extract(parsed, document_type=document_type, template_id="format_evidence")
    return parsed, result.fields


@pytest.mark.skipif(importlib.util.find_spec("pypdf") is None, reason="pypdf is not installed")
def test_parses_two_page_pdf_report(tmp_path: Path) -> None:
    sample = tmp_path / "report_evidence.pdf"
    document = canvas.Canvas(str(sample))
    document.drawString(72, 760, "Title: Multi-format Processing Report")
    document.drawString(72, 730, "Owner: qa.evidence")
    document.showPage()
    document.drawString(72, 760, "Date: 2026-08-10")
    document.drawString(72, 730, "Summary: PDF parsing remained traceable across two pages.")
    document.save()

    parsed, fields = extract_fields(sample)

    assert parsed.parser_name == "pypdf"
    assert parsed.metadata["page_count"] == 2
    assert fields["title"] == "Multi-format Processing Report"
    assert fields["owner"] == "qa.evidence"
    assert fields["date"] == "2026-08-10"


def test_parses_docx_contract_table(tmp_path: Path) -> None:
    sample = tmp_path / "contract_evidence.docx"
    document = Document()
    table = document.add_table(rows=2, cols=4)
    headers = ["合同编号", "甲方", "乙方", "生效日期"]
    values = ["DEMO-DOCX-001", "青岚示例有限公司", "方舟示例有限公司", "2026-08-10"]
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for index, value in enumerate(values):
        table.rows[1].cells[index].text = value
    document.save(sample)

    parsed, fields = extract_fields(sample)

    assert parsed.parser_name == "python-docx"
    assert parsed.metadata["table_count"] == 1
    assert fields["document_no"] == "DEMO-DOCX-001"
    assert fields["party_a"] == "青岚示例有限公司"
    assert fields["party_b"] == "方舟示例有限公司"


@pytest.mark.filterwarnings("ignore:datetime.datetime.utcnow.*:DeprecationWarning:openpyxl.*")
def test_parses_xlsx_invoice_table(tmp_path: Path) -> None:
    sample = tmp_path / "invoice_evidence.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Invoice"
    worksheet.append(["invoice number", "issuer", "date", "total"])
    worksheet.append(["DEMO-XLSX-001", "River Evidence Systems Ltd.", "2026-08-10", "8200.00"])
    workbook.save(sample)
    workbook.close()

    parsed, fields = extract_fields(sample)

    assert parsed.parser_name == "openpyxl"
    assert parsed.metadata["sheet_count"] == 1
    assert fields["document_no"] == "DEMO-XLSX-001"
    assert fields["issuer"] == "River Evidence Systems Ltd."
    assert fields["amount"] == "8200.00"


@pytest.mark.skipif(
    shutil.which("tesseract") is None or not OCR_FONT.is_file(),
    reason="Tesseract or the CI OCR test font is not installed",
)
def test_reads_high_contrast_invoice_image_with_tesseract(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    sample = tmp_path / "invoice_ocr_evidence.png"
    image = Image.new("RGB", (900, 160), "white")
    ImageDraw.Draw(image).text(
        (30, 48),
        "INVOICE DEMO OCR 004",
        fill="black",
        font=ImageFont.truetype(str(OCR_FONT), 48),
    )
    image.save(sample)

    monkeypatch.setenv("OCR_LANGUAGES", "eng")
    get_settings.cache_clear()
    try:
        parsed = DocumentParser().parse(sample, sample.name)
    finally:
        get_settings.cache_clear()

    assert parsed.parser_name == "tesseract"
    assert "DEMO OCR 004" in parsed.text.upper()
    assert parsed.metadata["image_size"] == [900, 160]
